# 这个小程序用来生成一份课堂演讲用的 PPT。
# 主题：从《寂静的春天》谈环境伦理，并结合身边的所见所想。
# 约束：正文大约 1000 字，总页数不超过 15 页。
#
# 运行方法（在终端里）：
#     python3 生成演讲PPT.py
# 运行后，同目录会出现“寂静的春天-环境伦理演讲.pptx”。

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


# 幻灯片用到的几种颜色。RGB 是红、绿、蓝三个数，每个数从 0 到 255。
# 这里选深绿、米色和陶土色，是为了让页面安静，适合谈自然，而不是像广告。
FOREST = RGBColor(0x1B, 0x3A, 0x2F)
CREAM = RGBColor(0xF6, 0xF3, 0xEC)
PAPER = RGBColor(0xFF, 0xFC, 0xF8)
INK = RGBColor(0x1C, 0x28, 0x24)
CLAY = RGBColor(0xA8, 0x5A, 0x3A)
MUTED = RGBColor(0x5C, 0x68, 0x60)
FOOT_TEXT = RGBColor(0xE4, 0xEB, 0xE4)

# 系统里已经安装的中文字体。必须写进文本框，否则中文可能变成方框。
FONT = "WenQuanYi Micro Hei"

# 宽屏比例 16:9。单位是英寸，后面用 Inches() 转成 PPT 内部单位。
SLIDE_W = 13.333
SLIDE_H = 7.5


# 下面这份列表，就是演讲稿本身。
# 每一项是一页：小标题（页眉）、大标题、两段正文。
# 先把话写在这里，再交给后面的函数“贴”到幻灯片上，改稿时只改这一处。
PAGES = [
    {
        "kicker": "从一本书读起",
        "title": "一种不正常的安静",
        "paras": [
            "春天本来喧哗。鸟、虫和风叠在一起，人才觉得这个季节是活的。",
            "卡森写下“寂静的春天”，说的是另一种安静：\n不是安宁，而是许多生命已经无法出声。\n安静若来得不正常，就该被追问。我由此想起的，是自己每天走过的地方。",
        ],
    },
    {
        "kicker": "书里的追问",
        "title": "后果并不听话",
        "paras": [
            "《寂静的春天》写到药剂怎样进入土壤、河水、食物和身体，\n又怎样沿着食物链慢慢积累。",
            "她并不要求人们放弃防治虫害。她反对的是另一种自信：\n以为伤害可以只落在指定的那一只虫子身上。",
        ],
    },
    {
        "kicker": "伦理要问的",
        "title": "有用之后，还要问",
        "paras": [
            "环境伦理不停止于“这样做有没有用”。\n它还要问：对谁有用，伤害了谁，谁没有同意，损失能不能挽回。",
            "鸟不再鸣叫，既是生态上的损失，也是道德上的事故。\n消失的声音，补偿并不总能换回。",
        ],
    },
    {
        "kicker": "几种看法",
        "title": "先守住底线",
        "paras": [
            "有人因为人类需要而保护自然，有人因为生命本身值得尊重，\n也有人看到河流、土壤与物种构成的共同体。这些名称可以先放下。",
            "即便只为人的健康，我们已经应当克制。\n不能开口的生命，更不该因为无法拒绝就被忽略。",
        ],
    },
    {
        "kicker": "身边所见",
        "title": "请关窗",
        "paras": [
            "夏夜，小区喷药灭蚊，通知常常只剩一句“请关窗”。\n草地、晾着的衣服、灌木里的猫，都躲不开那层雾。",
            "不想被叮咬，这个愿望没有错。我们很少接着问：\n还有没有更窄的办法？知情，是不是等于被通知了一句？",
        ],
    },
    {
        "kicker": "身边所见",
        "title": "整洁不等于活着",
        "paras": [
            "河岸被砌平，路灯亮了，老人却说鱼一年比一年少。\n整洁有时是治理，有时只是清除。",
            "外卖和塑料袋更日常：没有谁存心作恶，只有不断重复的方便。\n代价被推给河道、土壤，和还不能发言的后人。",
        ],
    },
    {
        "kicker": "责任如何消失",
        "title": "“不是我”",
        "paras": [
            "这些伤害看不见，来得慢，又被分成很小的一份，\n于是每个人都能说“不是我”。责任一旦被稀释，并不因此消失。",
            "环境问题里最常见的不公，是好处归一方，风险归另一方。\n卡森把散开的事实重新接上，让链条末端的生命被看见。",
        ],
    },
    {
        "kicker": "时间与后人",
        "title": "再等一等",
        "paras": [
            "“还没有定论”常常用来换取继续照旧的时间。\n当伤害可能严重、而且难以挽回，谨慎本身就是一种道德。",
            "过早的谨慎，往往比过晚的确信更负责。\n未来的人不坐在今天的谈判桌前。\n对他们来说，再等一等，也许就是没有了。",
        ],
    },
    {
        "kicker": "一种幻觉",
        "title": "何处该停",
        "paras": [
            "卡森质疑“控制自然”，不是反对求知，而是反对狂妄。\n认识了局部，不等于有权重组整个生活世界。",
            "全面喷药、砌死河岸、再换一个更大的袋子，\n每一步都像在解决问题，连起来却可能去掉自然自身的调节。",
        ],
    },
    {
        "kicker": "落到今天",
        "title": "在意，才会出声",
        "paras": [
            "责任不必先说得很大。问清一次喷药的范围，少用一个塑料袋，\n同时承认这代替不了制度。更要紧的是还愿意在意。",
            "没有在意，问题就不会被说出来。\n人不该把世界改到只剩下自己的回声。\n寂静若已开始，它不是风景，是催促。",
        ],
    },
]


def set_run_font(run, size, color, bold=False):
    """给一段文字设置字体、字号和颜色。

    PPT 里中文要用“东亚字体”(ea)。只设置西文字体时，
    换一台电脑打开，汉字有时会找不到字形。所以这里三种字体一起写上。
    """
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color

    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, qn(tag))
        node.set("typeface", FONT)


def add_rect(slide, x, y, w, h, color):
    """在页面上画一个没有边框的矩形，用来做色块背景。"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def write_paragraph(paragraph, text, size, color, bold, align, spacing):
    """把一段话写进段落。

    字符串里的换行符 \\n 会变成同一段里的手动换行。
    这样可以自己决定在哪里断行，避免“食物链”“拒绝”被拆到两行。
    """
    paragraph.alignment = align
    paragraph.line_spacing = spacing
    parts = text.split("\n")
    for index, part in enumerate(parts):
        run = paragraph.add_run()
        run.text = part
        set_run_font(run, size, color, bold)
        if index < len(parts) - 1:
            # a:br 是 PPT 里的换行标记，必须放在两个文字块之间。
            run._r.addnext(etree.Element(qn("a:br")))


def add_text(slide, x, y, w, h, lines, size, color, bold=False, align=PP_ALIGN.LEFT, spacing=1.15):
    """在指定位置放一个文本框。

    lines 可以是一行字，也可以是多段（列表）。
    spacing 是行距倍数，大于 1 会让字不挤在一起，方便朗读。
    返回文本框，方便后面再调整段距。
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.TOP

    # 文本框默认留白太大，演讲页会显得字被往里挤，这里把内边距改小。
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)

    if isinstance(lines, str):
        lines = [lines]

    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        write_paragraph(paragraph, line, size, color, bold, align, spacing)
    return frame


def add_notes(slide, text):
    """把这一页要念的话写进备注。放映时打开“演讲者视图”就能看见。"""
    notes = slide.notes_slide.notes_text_frame
    notes.text = text


def count_speech_chars():
    """统计正文汉字和标点，用来核对是不是大约 1000 字。

    封面和页脚不计入。标题和正文计入，因为上台时这两部分都会被讲到。
    """
    total = 0
    for page in PAGES:
        total += len(page["title"])
        for para in page["paras"]:
            total += len(para.replace("\n", ""))
    return total


def build_cover(slide):
    """封面：只交代题目和篇幅，不提前把正文讲完。"""
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    add_rect(slide, 0, 0, 0.22, SLIDE_H, FOREST)
    add_rect(slide, 0, 6.72, SLIDE_W, 0.78, FOREST)

    add_text(slide, 0.7, 1.85, 10, 0.4, "环境伦理  ·  读书报告", 16, CLAY, True)
    add_text(slide, 0.7, 2.4, 12, 1.2, "寂静从何而来", 54, FOREST, True)
    add_rect(slide, 0.78, 3.75, 2.1, 0.045, CLAY)
    add_text(
        slide,
        0.7,
        4.0,
        11,
        0.5,
        "重读蕾切尔·卡森《寂静的春天》",
        22,
        INK,
    )
    add_text(
        slide,
        0.7,
        4.7,
        10.5,
        0.8,
        "把书中的警告，放到身边那些被当成小事的安静里。",
        18,
        MUTED,
    )
    add_text(
        slide,
        0.7,
        6.86,
        8,
        0.42,
        "正文约一千字    ·    建议讲述五分钟",
        15,
        FOOT_TEXT,
    )
    add_text(
        slide,
        10.2,
        6.86,
        2.6,
        0.42,
        "共 11 页",
        15,
        FOOT_TEXT,
        align=PP_ALIGN.RIGHT,
    )
    add_notes(
        slide,
        "封面停十几秒即可。先念题目，再补一句：今天不复述整本书，只谈它怎样改变我看身边事物的方式。",
    )


def build_content(slide, page, index, total_pages):
    """正文页：上面是题目，中间是要念的两段话，下面是页码。"""
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, CREAM)
    add_rect(slide, 0, 0, SLIDE_W, 0.12, FOREST)
    add_rect(slide, 0, 7.08, SLIDE_W, 0.42, FOREST)

    # 右侧大号页码只做装饰，颜色很浅，不抢正文。
    add_text(
        slide,
        10.6,
        0.22,
        2.3,
        0.9,
        f"{index:02d}",
        32,
        RGBColor(0xD5, 0xDE, 0xD6),
        True,
        align=PP_ALIGN.RIGHT,
    )

    add_text(slide, 0.62, 0.38, 8, 0.36, page["kicker"], 14, CLAY, True)
    add_text(slide, 0.58, 0.82, 10.5, 0.85, page["title"], 36, FOREST, True)
    add_rect(slide, 0.7, 1.78, 1.35, 0.045, CLAY)

    # 白底卡片只托住正文，不要拉得太高，否则下面会空出一大块。
    add_rect(slide, 0.55, 2.12, 12.25, 4.55, PAPER)
    add_rect(slide, 0.55, 2.12, 0.08, 4.55, CLAY)

    body = add_text(
        slide,
        0.85,
        2.38,
        11.55,
        4.05,
        page["paras"],
        24,
        INK,
        spacing=1.28,
    )
    # 只在两段之间留空，最后一段下面不再多垫一行。
    for paragraph in body.paragraphs[:-1]:
        paragraph.space_after = Pt(16)

    add_text(
        slide,
        0.55,
        7.12,
        8,
        0.32,
        "《寂静的春天》  ·  环境伦理",
        12,
        FOOT_TEXT,
    )
    add_text(
        slide,
        10.3,
        7.12,
        2.5,
        0.32,
        f"{index}  /  {total_pages}",
        12,
        FOOT_TEXT,
        align=PP_ALIGN.RIGHT,
    )

    spoken = page["title"] + "。" + "".join(page["paras"])
    add_notes(
        slide,
        "建议用时约三十秒。可以基本照着念，念完停一下再翻页。\n\n" + spoken,
    )


def main():
    speech_chars = count_speech_chars()
    page_count = 1 + len(PAGES)  # 1 是封面
    if page_count > 15:
        raise SystemExit("页数超过 15 页，请删减后再生成。")

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "寂静从何而来：重读《寂静的春天》"
    prs.core_properties.subject = "环境伦理读书报告"
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    build_cover(cover)

    for index, page in enumerate(PAGES, start=2):
        slide = prs.slides.add_slide(blank)
        build_content(slide, page, index, page_count)

    output = "寂静的春天-环境伦理演讲.pptx"
    prs.save(output)

    # 打印出来，方便核对作业要求：大约 1000 字、不超过 15 页。
    print(f"已保存：{output}")
    print(f"总页数：{page_count}")
    print(f"标题加正文：{speech_chars} 字")


if __name__ == "__main__":
    main()
